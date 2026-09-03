#!/usr/bin/env python3
"""
Gradient-based reframing of the behavioural and companion-cohort questions.

WHY THIS EXISTS
---------------
Two colleagues who read the draft made the same point independently, and it is correct:
the manuscript's behavioural and addiction analyses ask the wrong question.

The study's finding is not that visual cortex has some MEAN level of
controllability. It is that controllability varies systematically WITH
RETINOTOPIC POSITION -- a per-participant gradient. Correlating a
participant's mean controllability with delay discounting, or comparing mean
controllability between cohorts, tests a quantity the paper never claimed was
organised. It is unsurprising that both came back null.

The question that matches the finding is whether the STRENGTH OF THE GRADIENT
differs between people, and whether that difference tracks behaviour or
substance-use exposure.

WHAT THIS DOES
--------------
(1) Extracts, per participant, the eccentricity-controllability partial
    correlation -- the per-participant gradient strength, Fisher-z scaled.
(2) Correlates that gradient against each behavioural measure.
(3) Tests a cohort x eccentricity interaction: does the gradient differ
    between the primary and companion cohorts?
(4) Within the companion cohort, tests a severity x eccentricity interaction.

(3) and (4) are group-level tests on per-participant gradients, so the unit of
inference is the participant throughout, and the small-spatial-unit
chance-alignment problem documented elsewhere in this project does not apply.
Significance uses permutation (label shuffling) rather than parametric tests,
matching project convention.

HONEST FRAMING, STATED IN ADVANCE
---------------------------------
The companion cohort is N=20. These interaction tests are better-posed than
the mean-difference tests they replace, but they are not better-powered. A
null here should be reported as "the gradient was not detectably altered",
not as evidence that substance use leaves the gradient intact.

Usage:
    python3 run_gradient_analyses.py \
        --prf-backbone-csv all_subjects_prf_backbone.csv \
        --controllability-csv "Received files/all_subjects_controllability_glasser.csv" \
        --atlas glasser --master-table master-table.xlsx \
        --out-prefix gradient_glasser
    python3 run_gradient_analyses.py --self-test
"""
import argparse, csv, math, re, sys
from collections import defaultdict

import numpy as np
from scipy import stats

GLASSER_VIS_AREAS = {"V1","V2","V3","V4","V3A","V3B","V6","V6A","V7","IPS1",
                     "FFC","PIT","V8","VMV1","VMV2","VMV3","VVC",
                     "FST","LO1","LO2","LO3","MST","MT","PH","V3CD","V4t"}
GLASSER_RE = re.compile(r"^[LR]_(.+)_ROI$")
BEHAV_COLS = ["Delay Discounting AUC ($200)", "Delay Discounting AUC ($40,000)",
              "Flanker Inhibitory Control (age-adj)"]
SEVERITY_COLS = ["Alcohol Abuse Symptom Count", "Alcohol Dependence Symptom Count",
                 "Total Drinks, Past 7 Days", "Marijuana Times-Used Category",
                 "Lifetime Smoking History Class"]


def is_vis(lab):
    m = GLASSER_RE.match(lab)
    return bool(m) and m.group(1) in GLASSER_VIS_AREAS


def load_prf(path, minv, atlas):
    ecc, size = defaultdict(dict), {}
    with open(path, newline="") as f:
        for r in csv.DictReader(f):
            if r["atlas"] != atlas: continue
            v = (r["is_vis_network"]=="True") if atlas=="schaefer200" else is_vis(r["backbone_label"])
            if not v: continue
            size[r["backbone_label"]] = int(r["n_vertices_total"])
            if r["mean_eccentricity_deg"]=="" : continue
            if int(r["n_vertices_valid_R2gt30"]) < minv: continue
            ecc[r["subject"]][r["backbone_label"]] = float(r["mean_eccentricity_deg"])
    return ecc, size


def load_ctrl(path):
    avg, mod = defaultdict(dict), defaultdict(dict)
    with open(path, newline="") as f:
        for r in csv.DictReader(f):
            a = float(r["average_controllability"])
            mod[r["subject"]][r["node_label"]] = float(r["modal_controllability"])
            if a > 0: avg[r["subject"]][r["node_label"]] = math.log(a)
    return avg, mod


def partial_corr(x,y,z):
    x,y,z = map(lambda v: np.asarray(v,float), (x,y,z))
    if np.std(x)==0 or np.std(y)==0 or np.std(z)==0: return np.nan
    rxy,rxz,ryz = stats.pearsonr(x,y)[0], stats.pearsonr(x,z)[0], stats.pearsonr(y,z)[0]
    den = math.sqrt((1-rxz**2)*(1-ryz**2))
    return np.nan if den==0 or not np.isfinite(den) else (rxy-rxz*ryz)/den


def bh_fdr(p):
    p=np.asarray(p,float); n=len(p); o=np.argsort(p)
    q=p[o]*n/(np.arange(n)+1); q=np.minimum.accumulate(q[::-1])[::-1]
    out=np.empty(n); out[o]=np.clip(q,0,1); return out


def fz(r): return np.arctanh(np.clip(r,-0.999999,0.999999))


def per_subject_gradient(ecc, metric, size, min_nodes):
    """{subject: Fisher-z gradient strength}"""
    out={}
    for s,e in ecc.items():
        m=metric.get(s,{})
        common=sorted(set(e)&set(m)&set(size))
        if len(common)<min_nodes: continue
        r=partial_corr([e[n] for n in common],[m[n] for n in common],[size[n] for n in common])
        if np.isfinite(r): out[s]=float(fz(r))
    return out


def perm_corr(x, y, n_perm, rng):
    """Pearson r with a label-shuffling permutation p-value."""
    x,y=np.asarray(x,float),np.asarray(y,float)
    if len(x)<5 or np.std(x)==0 or np.std(y)==0: return np.nan, np.nan, len(x)
    obs=stats.pearsonr(x,y)[0]
    null=np.array([stats.pearsonr(rng.permutation(x),y)[0] for _ in range(n_perm)])
    p=(np.sum(np.abs(null)>=abs(obs))+1)/(n_perm+1)
    return float(obs), float(p), len(x)


def perm_group_diff(a, b, n_perm, rng):
    """Difference in mean gradient between two groups, permutation-tested."""
    a,b=np.asarray(a,float),np.asarray(b,float)
    obs=a.mean()-b.mean()
    pool=np.concatenate([a,b]); na=len(a)
    null=np.empty(n_perm)
    for i in range(n_perm):
        p=rng.permutation(pool); null[i]=p[:na].mean()-p[na:].mean()
    return float(obs), float((np.sum(np.abs(null)>=abs(obs))+1)/(n_perm+1))


def self_test():
    rng=np.random.default_rng(0)
    print("1. gradient extraction recovers a planted per-subject gradient")
    size={f"n{i}":100+i for i in range(20)}; nodes=list(size)
    ecc,met={},{}
    strengths=rng.uniform(0.2,0.9,30)
    for i,s in enumerate(map(str,range(30))):
        e=rng.normal(0,1,20); ecc[s]=dict(zip(nodes,e))
        met[s]=dict(zip(nodes, strengths[i]*e + rng.normal(0,0.4,20)))
    g=per_subject_gradient(ecc,met,size,15)
    r=stats.pearsonr([g[s] for s in map(str,range(30))], strengths)[0]
    assert r>0.6, r
    print(f"   ok (recovered gradient correlates r={r:.2f} with planted strength)")

    print("2. gradient-behaviour test detects a planted association")
    beh=strengths*3 + rng.normal(0,0.3,30)
    obs,p,n=perm_corr([g[s] for s in map(str,range(30))], beh, 2000, rng)
    assert p<0.01, (obs,p)
    print(f"   ok (r={obs:.2f}, p={p:.4f})")

    print("3. gradient-behaviour test is null for unrelated behaviour")
    obs,p,n=perm_corr([g[s] for s in map(str,range(30))], rng.normal(0,1,30), 2000, rng)
    assert p>0.05, (obs,p)
    print(f"   ok (r={obs:+.2f}, p={p:.3f})")

    print("4. group difference detects a planted shift")
    obs,p=perm_group_diff(rng.normal(0.6,0.2,25), rng.normal(0.2,0.2,20), 2000, rng)
    assert p<0.01 and obs>0.2, (obs,p)
    print(f"   ok (diff={obs:.2f}, p={p:.4f})")

    print("5. group difference is null for identical groups")
    obs,p=perm_group_diff(rng.normal(0.4,0.2,25), rng.normal(0.4,0.2,20), 2000, rng)
    assert p>0.05, (obs,p)
    print(f"   ok (diff={obs:+.3f}, p={p:.3f})")
    print("\nSELF-TEST: ALL PASSED")


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("--self-test",action="store_true")
    ap.add_argument("--prf-backbone-csv"); ap.add_argument("--controllability-csv")
    ap.add_argument("--atlas",default="glasser",choices=("schaefer200","glasser"))
    ap.add_argument("--master-table")
    ap.add_argument("--companion-prf-csv",default=None)
    ap.add_argument("--companion-controllability-csv",default=None)
    ap.add_argument("--companion-cohort-xlsx",default=None,
                    help="addiction-cohort-selection.xlsx -- enables the severity x eccentricity "
                         "interaction test within the companion cohort")
    ap.add_argument("--out-prefix")
    ap.add_argument("--min-valid-nodes",type=int,default=15)
    ap.add_argument("--min-vertices-per-node",type=int,default=5)
    ap.add_argument("--n-permutations",type=int,default=10000)
    ap.add_argument("--seed",type=int,default=42)
    a=ap.parse_args()
    if a.self_test: self_test(); return
    for r in ("prf_backbone_csv","controllability_csv","out_prefix"):
        if not getattr(a,r): ap.error(f"--{r.replace('_','-')} is required")

    rng=np.random.default_rng(a.seed)
    ecc,size = load_prf(a.prf_backbone_csv,a.min_vertices_per_node,a.atlas)
    avg,mod  = load_ctrl(a.controllability_csv)
    grads = {"average_controllability_log": per_subject_gradient(ecc,avg,size,a.min_valid_nodes),
             "modal_controllability":       per_subject_gradient(ecc,mod,size,a.min_valid_nodes)}
    for k,v in grads.items():
        arr=np.array(list(v.values()))
        print(f"{k}: n={len(v)} gradient mean z={arr.mean():+.3f} sd={arr.std():.3f} "
              f"range [{arr.min():+.3f},{arr.max():+.3f}]")

    rows=[]
    # ---- (2) gradient vs behaviour ----
    if a.master_table:
        import openpyxl
        wb=openpyxl.load_workbook(a.master_table,read_only=True); ws=wb["Subjects"]
        hdr=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
        idx={h:i for i,h in enumerate(hdr)}
        beh=defaultdict(dict)
        for row in ws.iter_rows(min_row=2,values_only=True):
            sid=str(int(row[idx["Subject ID"]])) if row[idx["Subject ID"]] is not None else None
            if not sid: continue
            for c in BEHAV_COLS:
                if c in idx and isinstance(row[idx[c]],(int,float)):
                    beh[c][sid]=float(row[idx[c]])
        for metric,g in grads.items():
            for c in BEHAV_COLS:
                shared=sorted(set(g)&set(beh.get(c,{})))
                if len(shared)<5: continue
                r,p,n=perm_corr([g[s] for s in shared],[beh[c][s] for s in shared],a.n_permutations,rng)
                rows.append({"analysis":"gradient_vs_behaviour","metric":metric,"variable":c,
                             "n":n,"estimate":r,"perm_p_value":p})

    # ---- (3) cohort x eccentricity interaction ----
    if a.companion_prf_csv and a.companion_controllability_csv:
        cecc,csize = load_prf(a.companion_prf_csv,a.min_vertices_per_node,a.atlas)
        cavg,cmod  = load_ctrl(a.companion_controllability_csv)
        cgr = {"average_controllability_log": per_subject_gradient(cecc,cavg,csize,a.min_valid_nodes),
               "modal_controllability":       per_subject_gradient(cecc,cmod,csize,a.min_valid_nodes)}
        for metric in grads:
            pa=list(grads[metric].values()); pb=list(cgr[metric].values())
            if len(pb)<5: continue
            obs,p=perm_group_diff(pa,pb,a.n_permutations,rng)
            rows.append({"analysis":"cohort_x_eccentricity","metric":metric,
                         "variable":"primary_minus_companion","n":len(pa)+len(pb),
                         "estimate":obs,"perm_p_value":p})

    # ---- (4) severity x eccentricity WITHIN the companion cohort ----
    # The biologically relevant question is not whether mean
    # controllability differs by severity, but whether severity alters the
    # retinotopic GRADIENT itself. Tested here as a correlation between each
    # companion participant's gradient strength and their severity score.
    if a.companion_cohort_xlsx and a.companion_prf_csv and a.companion_controllability_csv:
        import openpyxl
        wb = openpyxl.load_workbook(a.companion_cohort_xlsx, read_only=True)
        ws = wb["Cohort"]
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        idx = {h: i for i, h in enumerate(hdr)}
        sev = defaultdict(dict)
        for row in ws.iter_rows(min_row=2, values_only=True):
            raw = row[idx["Subject ID"]]
            if raw is None: continue
            sid = str(int(raw))
            for c in SEVERITY_COLS:
                if c in idx and isinstance(row[idx[c]], (int, float)):
                    sev[c][sid] = float(row[idx[c]])
        for metric, g in cgr.items():
            for c in SEVERITY_COLS:
                shared = sorted(set(g) & set(sev.get(c, {})))
                if len(shared) < 5: continue
                r, p, n = perm_corr([g[s] for s in shared], [sev[c][s] for s in shared],
                                    a.n_permutations, rng)
                rows.append({"analysis": "severity_x_eccentricity", "metric": metric,
                             "variable": c, "n": n, "estimate": r, "perm_p_value": p})

    if not rows:
        sys.exit("No analyses ran -- supply --master-table and/or companion inputs.")
    qs=bh_fdr([r["perm_p_value"] for r in rows])
    for r,q in zip(rows,qs): r["q_value_fdr"]=float(q)

    print("\n=== results ===")
    for r in sorted(rows,key=lambda x:x["perm_p_value"]):
        print(f"  {r['analysis']:24s} {r['metric']:28s} {r['variable'][:34]:34s} "
              f"n={r['n']:3d} est={r['estimate']:+.3f} p={r['perm_p_value']:.4f} q={r['q_value_fdr']:.4f}")

    out=f"{a.out_prefix}.csv"
    fields=["analysis","metric","variable","n","estimate","perm_p_value","q_value_fdr"]
    with open(out,"w",newline="") as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader()
        for r in rows: w.writerow({k:r.get(k,"") for k in fields})
    print(f"\nWritten to {out}")


if __name__=="__main__":
    main()
